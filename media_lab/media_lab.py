"""media_lab - point it at a folder; it transcribes every video/audio file inside it (recursively).

Per file -> media_lab/out/<folder name>/<relative path>/
    transcript.txt       clean text, paragraphs
    transcript_ts.txt    [hh:mm:ss] line per segment
    segments.json        raw whisper segments
    meta.json            path, size, duration, words, wpm, timings, visual stats
  video only:
    frames/*.jpg         SMART frames (see below), file names carry the timestamp
    sheet_01.jpg ...     contact sheets, 20 frames each, labelled hh:mm:ss
    ocr.txt              slide / on-screen text per frame (GPU, after whisper; slide region only)
    outline.md           long videos: timeline of "what was on screen" + "what was said" per visual section
  learning materials (every other file: pdf, docx, pptx, xlsx, txt, md, html, epub, images, subtitles...):
    copied 1:1 to the same relative path  + "<name>.extracted.txt" with the text (pdf/docx/pptx/xlsx/html/epub;
    images and scanned PDF pages are OCR'd on CPU)
Per folder -> out/<folder name>/_all_transcripts.md (outlines: speech + on-screen text), _all_materials.md, _manifest.md
            and out/<folder name>.zip (the whole output folder; rebuilt after every run, old zip removed first)
Uploads in progress: files changed in the last 2 min are "waiting"; the worker keeps rescanning (GPU released
while it waits) and picks them up when the upload is complete. Partly-decodable media are rejected, not marked done.

Smart frames
* short videos (<= 3 min): uniform, one frame every ~10 s (like reel_lab)
* long videos: a fast scan of keyframes only (tiny 128x72 greyscale, CPU, no full decode) finds where the
  picture really changes (new slide, screen share, camera switch). One frame per visual section, picked late
  in the section (slide fully built). Same-looking frames (speaker view, a slide shown again) are dropped.
  Gaps longer than 5 min get a check frame (also de-duplicated). Budget ~1.5 frames per minute, max 150.
  -> a 2 h talking-head call ends with a handful of frames; a 2 h slide webinar with one frame per slide.
* scan + frames run on CPU in parallel with whisper; OCR runs on the GPU right after whisper on the same file
  (slide region only, duplicate frames dropped first).

Safety
* source files are only READ - never moved, renamed, modified or deleted; nothing is written next to them
* symlinked sub-folders are NOT followed (e.g. lab/models)
* GPU: RTX 4090 only, 25% VRAM cap (reel_lab settings), shares reel_lab's gpu.lock = one GPU job at a time
* worker runs with low CPU priority (nice 10) and a limited number of CPU threads
* runs as a detached child process; on exit (done, error or stop()) it frees VRAM and wipes media_lab/tmp
* re-running the same folder skips files already done (same size + mtime) -> safe to resume
"""
import os, sys, json, time, glob, re, signal, subprocess
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "reel_lab"))
import reel_lab as rl  # noqa: E402

OUT, TMP, LOGS, JOBS = HERE / "out", HERE / "tmp", HERE / "logs", HERE / "jobs"
INDEX, STATE = HERE / "index.jsonl", HERE / "state.json"
for _p in (OUT, TMP, LOGS, JOBS):
    _p.mkdir(parents=True, exist_ok=True)

VIDEO = {".mp4", ".mov", ".mkv", ".avi", ".webm", ".m4v", ".mpg", ".mpeg", ".wmv", ".flv", ".3gp", ".mts", ".ts"}
AUDIO = {".mp3", ".wav", ".m4a", ".aac", ".flac", ".ogg", ".opus", ".wma", ".aiff", ".aif", ".amr"}
BATCH = 4            # whisper batch; falls back to 1 on OOM (still inside the 25% cap)
SHORT_MAX = 180      # <= 3 min: uniform frames
SHORT_FRAMES = 60
LONG_PER_MIN = 1.5   # frame budget for long videos
LONG_MAX = 150
MIN_SECTION = 4.0    # a visual section shorter than this is a transition, ignore
COVER_GAP = 300      # no gap between frames longer than 5 min
SCAN_W, SCAN_H = 128, 72
PIX_DIFF = 18        # grey-level change that counts as a changed pixel
CUT = 0.01           # >1% of still pixels changed = new visual state
FRAME_W = 960        # readable slides for OCR / AI
CPU_THREADS = max(2, min(8, (os.cpu_count() or 8) // 4))
SKIP_DIRS = {HERE.resolve(), rl.ROOT.resolve()}
SETTLE_S = 120       # a file modified less than 2 min ago is still being uploaded
WAIT_MAX_S = 4 * 3600  # keep waiting for uploads at most 4 h after the last change
RES_MAX = 500 * 1024 ** 2  # don't copy non-media files bigger than 500 MB (listed in the manifest instead)
PDF_OCR_PAGES = 80   # OCR at most this many text-less PDF pages per file
VENDOR = HERE / "vendor"   # private pip target for pymupdf / python-docx / python-pptx / openpyxl
TEXTLIKE = {".txt", ".md", ".csv", ".json", ".srt", ".vtt", ".tsv", ".xml", ".yaml", ".yml"}
IMAGES = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif"}
OCR_LANGS = ["en"]  # EasyOCR languages, e.g. ["en", "uk"] or ["en", "ru"] (one non-Latin script per reader)
OCR_GPU = True  # OCR runs on the GPU right after whisper finished the file (5-10x faster than CPU)
DHASH_MAX_DIST = 12  # frames whose slide-region dHash differs by <= this are the same picture -> dropped
SPEECH_MIN_S = 1.0  # a piece needs this much speech-level audio to be transcribed at all (kills ghost text)
HALLUCINATIONS = re.compile(r"^(thank you\.?|thanks for watching\.?|subtitles? by .*|please subscribe.*|"
                            r"like and subscribe.*|you're back,? you're back.*|\.+|…)$", re.I)


# ---------------- helpers ----------------
def _hms(t):
    t = int(t or 0)
    return f"{t // 3600:02d}:{t % 3600 // 60:02d}:{t % 60:02d}"


def _walk(folder):
    folder = Path(folder).expanduser().resolve()
    if not folder.is_dir():
        raise FileNotFoundError(f"not a folder: {folder}")
    media, res = [], []
    for dp, dn, fn in os.walk(folder, followlinks=False):
        dn[:] = sorted(d for d in dn if not d.startswith(".") and (Path(dp) / d).resolve() not in SKIP_DIRS
                       and not (Path(dp) / d).is_symlink())
        for f in sorted(fn):
            p = Path(dp) / f
            if f.startswith(".") or f.startswith("~$") or p.is_symlink() or f in ("Thumbs.db", "desktop.ini"):
                continue
            (media if p.suffix.lower() in VIDEO | AUDIO else res).append(p)
    return folder, media, res


def find_media(folder):
    folder, media, _ = _walk(folder)
    return folder, media


def _uploading(p):
    """True while a file is still being written (modified in the last SETTLE_S seconds)."""
    try:
        return time.time() - p.stat().st_mtime < SETTLE_S
    except FileNotFoundError:
        return True


def out_dir(folder, src):
    return OUT / Path(folder).name / Path(src).relative_to(folder)


def _duration(p):
    try:
        return rl._probe(p)["duration"]
    except Exception:
        return 0.0


def preview(folder):
    """What would be processed (no GPU, no writes)."""
    folder, files, res = _walk(folder)
    gb = sum(f.stat().st_size for f in files) / 1e9
    hours = sum(_duration(f) for f in files) / 3600
    done = sum(1 for f in files if _is_done(folder, f))
    up = [f for f in files + res if _uploading(f)]
    lines = [f"{folder}",
             f"  media: {len(files)} files, {gb:.2f} GB, {hours:.1f} h, {done} already done",
             f"  learning materials: {len(res)} files (copied + text extracted)",
             f"  still uploading right now: {len(up)}",
             f"  rough time: ~{hours / 20:.1f}-{hours / 12:.1f} h (whisper runs ~12-20x faster than real time)"]
    lines += ["  " + str(f.relative_to(folder)) + ("   <- uploading" if f in up else "") for f in (files + res)[:60]]
    if len(files) + len(res) > 60:
        lines.append(f"  ... +{len(files) + len(res) - 60} more")
    print("\n".join(lines))


def _is_done(folder, src):
    m = out_dir(folder, src) / "meta.json"
    if not m.exists():
        return False
    try:
        d = json.loads(m.read_text())
        st = src.stat()
        if not (d.get("ok") and d.get("size") == st.st_size and d.get("mtime") == int(st.st_mtime)):
            return False
        # results made before the completeness check: redo suspicious near-empty long files
        if "audio_seconds" not in d and (d.get("duration") or 0) > 600 and (d.get("wpm") or 0) < 10:
            return False
        return True
    except Exception:
        return False


def _state(**kw):
    d = {}
    if STATE.exists():
        try:
            d = json.loads(STATE.read_text())
        except Exception:
            pass
    d.update(kw, updated=time.strftime("%Y-%m-%d %H:%M:%S"))
    STATE.write_text(json.dumps(d, indent=1))
    return d


def _alive(pid):
    try:
        os.kill(int(pid), 0)
        return True
    except ProcessLookupError:
        return False
    except (PermissionError, ValueError, TypeError):
        return bool(pid)


def _is_our_worker(pid):
    try:
        return "media_lab.py" in Path(f"/proc/{int(pid)}/cmdline").read_bytes().decode(errors="ignore")
    except Exception:
        return False


# ---------------- speech (GPU, child process only) ----------------
SR = 16000
PIECE_MAX, PIECE_MIN = 28.0, 8.0     # whisper sees <= 30 s at a time; cut at the quietest moment
SILENCE_RMS = 0.004                  # pieces quieter than this are skipped (no hallucinated text)
ASR_GROUP = 64                       # pieces per pipeline call


def _load_audio(src):
    """Decode to 16 kHz mono. Damaged packets (common at the start of Zoom recordings) are skipped
    instead of failing the whole file; fails only if (almost) no audio comes out."""
    import numpy as np
    r = subprocess.run(["ffmpeg", "-nostdin", "-v", "error", "-err_detect", "ignore_err", "-max_error_rate", "1.0",
                        "-i", str(src), "-vn", "-ac", "1", "-ar", str(SR), "-f", "f32le", "pipe:1"],
                       capture_output=True)
    x = np.frombuffer(r.stdout, np.float32)
    err = r.stderr.decode(errors="ignore")
    if len(x) < SR * 5:
        raise RuntimeError("audio could not be decoded: " + err[-400:])
    bad = err.count("Invalid data found")
    if bad:
        print(f"   note: skipped {bad} damaged audio packets", flush=True)
    return x


def _pieces(x):
    """Split audio at natural pauses into <= 28 s pieces. Returns [(start_s, end_s)].
    A piece is kept only if it holds >= SPEECH_MIN_S of speech-level audio (a click or a cough is not speech)."""
    import numpy as np
    hop = int(0.03 * SR)
    n = len(x) // hop
    if n == 0:
        return []
    rms = np.sqrt((x[:n * hop].reshape(n, hop) ** 2).mean(axis=1))
    rms = np.convolve(rms, np.ones(10) / 10, mode="same")  # ~0.3 s smoothing: find real pauses
    out, cur = [], 0
    lo, hi = int(PIECE_MIN / 0.03), int(PIECE_MAX / 0.03)
    need = int(SPEECH_MIN_S / 0.03)
    while cur < n:
        if n - cur <= hi:
            end = n
        else:
            end = cur + lo + int(np.argmin(rms[cur + lo:cur + hi]))
        if int((rms[cur:end] >= SILENCE_RMS).sum()) >= need:
            out.append((cur * 0.03, min(end * 0.03, len(x) / SR)))
        cur = end
    return out


class Incomplete(RuntimeError):
    """The file can't be fully decoded yet (still uploading, or damaged)."""


def _check_complete(x, duration):
    got = len(x) / SR
    if duration > 60 and got < 0.85 * duration:
        raise Incomplete(f"incomplete file: only {_hms(got)} of {_hms(duration)} could be decoded "
                         f"(still uploading, or damaged - re-copy it)")
    return got


def _transcribe(x):
    """16 kHz audio -> [{start, end, text}] with real timestamps, batched on the GPU."""
    import torch
    pcs = _pieces(x)
    asr = rl._get_asr()
    segs = []
    with torch.inference_mode():
        for g in range(0, len(pcs), ASR_GROUP):
            grp = pcs[g:g + ASR_GROUP]
            inputs = [x[int(s * SR):int(e * SR)].copy() for s, e in grp]      # 16 kHz = whisper's own rate
            for bs in (BATCH, 1):
                try:
                    res = asr(inputs, batch_size=bs, return_timestamps=True,   # timestamps = fewer skipped words
                              generate_kwargs={"task": "transcribe"})
                    break
                except torch.cuda.OutOfMemoryError:
                    torch.cuda.empty_cache()
                    if bs == 1:
                        raise
            for (s, e), r in zip(grp, res):
                chunks = [c for c in (r.get("chunks") or []) if (c.get("text") or "").strip()]
                ok = chunks and all(c.get("timestamp") and c["timestamp"][0] is not None for c in chunks)
                if not ok:
                    segs.append({"start": round(s, 2), "end": round(e, 2), "text": (r.get("text") or "").strip()})
                    continue
                for c in chunks:                                    # sentence-level times inside the piece
                    c0, c1 = c["timestamp"]
                    c1 = c1 if c1 is not None else e - s
                    segs.append({"start": round(min(s + c0, e), 2), "end": round(min(s + c1, e), 2),
                                 "text": c["text"].strip()})
    return _clean(segs)


def _clean(segs):
    """Drop empty lines, whisper's repeated-line hallucinations (2nd identical line in a row) and classic ghosts."""
    out, prev = [], None
    for s in segs:
        t = s["text"].strip()
        if not t or HALLUCINATIONS.match(t):
            continue
        if t == prev:
            continue
        prev = t
        s["text"] = t
        out.append(s)
    return out


def _paragraphs(segs, gap=2.0, max_len=900):
    paras, cur, last_end = [], "", None
    for s in segs:
        brk = last_end is not None and s["start"] is not None and s["start"] - last_end > gap
        if cur and (brk or len(cur) > max_len):
            paras.append(cur.strip()); cur = ""
        cur += " " + s["text"]
        last_end = s["end"]
    if cur.strip():
        paras.append(cur.strip())
    return "\n\n".join(paras)


# ---------------- visuals (CPU) ----------------
def _scan(src, duration):
    """Tiny greyscale thumbnails through the whole video, cheap.
    Keyframes only (no full decode); if the file has too few keyframes (e.g. some Zoom exports),
    fall back to one sample every 4 s. Returns times[], thumbs (N,H,W) uint8, mode."""
    import numpy as np

    def grab(pre, vf):
        for vs in (["-fps_mode", "vfr"], ["-vsync", "vfr"]):     # new / old ffmpeg
            cmd = ["ffmpeg", "-nostdin", "-hide_banner", "-loglevel", "info", *pre, "-i", str(src), "-an", "-sn",
                   "-vf", f"{vf}scale={SCAN_W}:{SCAN_H},format=gray,showinfo", *vs,
                   "-f", "rawvideo", "pipe:1"]
            r = subprocess.run(cmd, capture_output=True)
            if r.returncode == 0:
                break
        ts = [float(x) for x in re.findall(rb"pts_time:\s*([\d.]+)", r.stderr)]
        n = len(r.stdout) // (SCAN_W * SCAN_H)
        k = min(n, len(ts))
        a = np.frombuffer(r.stdout[:k * SCAN_W * SCAN_H], np.uint8).reshape(k, SCAN_H, SCAN_W)
        return ts[:k], a

    ts, a = grab(["-skip_frame", "nokey"], "")
    mode = "keyframes"
    if len(ts) < duration / 12:          # keyframes sparser than every 12 s -> sample instead
        ts, a = grab([], "fps=1/4,")
        mode = "sampled-4s"
    return ts, a, mode


def _pick_long(ts, a, duration):
    """Choose frame times for a long video from the scan. Returns [(t, section_start, reason)], stats.
    Works on the share of pixels that changed, ignoring 'live' pixels that move almost all the time
    (webcam tile, speaker), so a new slide on the same template is still a cut and a moving speaker is not."""
    import numpy as np
    n = len(ts)
    if n < 3:
        return [(min(5.0, duration / 2), 0.0, "only")], {"scan_points": n}
    f = a.astype(np.int16)
    ch = np.abs(np.diff(f, axis=0)) > PIX_DIFF                    # (n-1, H, W) changed pixels
    live = ch.mean(axis=0) > 0.5                                   # moving most of the time
    still = ~live
    p = (ch & still).sum(axis=(1, 2)) / max(1, int(still.sum()))   # share of still pixels changed, per pair
    cut_thr = max(CUT, 2.5 * float(np.percentile(p, 20)))         # above the compression-noise floor
    cut = p > cut_thr
    dyn = np.zeros(n, bool)                                        # sample differs from both neighbours = motion
    dyn[1:-1] = cut[:-1] & cut[1:]
    dyn[0], dyn[-1] = cut[0], cut[-1]
    starts = [0]
    for i in range(1, n):
        if cut[i - 1] and not (dyn[i - 1] and dyn[i]):
            starts.append(i)
    starts.append(n)
    small = np.stack([_blur16(x) for x in f])
    sections = []
    for s, e in zip(starts[:-1], starts[1:]):
        t0 = ts[s]
        t1 = ts[e] if e < n else duration
        if t1 - t0 < MIN_SECTION:
            continue
        j = min(e - 1, s + int((e - s) * 0.8))                     # late in the section: slide fully built
        sections.append({"t": min(ts[j] + 1.0, t1 - 0.5), "start": t0, "idx": j, "len": t1 - t0,
                         "dyn": bool(dyn[s:e].mean() > 0.5)})
    budget = int(min(LONG_MAX, max(12, duration / 60 * LONG_PER_MIN)))
    cover, prev = [], 0.0
    for t in [x["t"] for x in sections] + [duration]:
        g = t - prev
        if g > COVER_GAP:
            k = int(g // COVER_GAP)
            for m in range(1, k + 1):
                tc = prev + g * m / (k + 1)
                j = int(np.argmin(np.abs(np.asarray(ts) - tc)))
                cover.append({"t": tc, "start": tc, "idx": j, "len": 0, "dyn": bool(dyn[j]), "cover": True})
        prev = t
    kept = []
    for c in sorted(sections, key=lambda x: -x["len"]) + cover:   # longest sections win the budget
        if len(kept) >= budget + len(cover):
            break
        dup = False
        for k in kept:
            if c["dyn"] and k["dyn"]:
                dup = np.abs(small[c["idx"]] - small[k["idx"]]).mean() < 10     # same speaker shot
            else:
                q = ((np.abs(f[c["idx"]] - f[k["idx"]]) > PIX_DIFF) & still).sum() / max(1, int(still.sum()))
                dup = q < cut_thr                                                  # same slide again
            if dup:
                break
        if not dup:
            kept.append(c)
    kept.sort(key=lambda x: x["t"])
    stats = {"scan_points": n, "live_px": round(float(live.mean()), 3), "cut": round(cut_thr, 3), "sections": len(sections),
             "cover_checks": len(cover), "budget": budget, "kept": len(kept)}
    stats["_live"] = live
    return [(c["t"], c["start"], "check" if c.get("cover") else ("motion" if c["dyn"] else "still")) for c in kept], stats


def _blur16(x):
    import cv2
    return cv2.resize(x.astype("float32"), (16, 9), interpolation=cv2.INTER_AREA)


def _grab(src, t, f, width=FRAME_W):
    subprocess.run(["ffmpeg", "-nostdin", "-v", "error", "-y", "-ss", f"{t:.2f}", "-i", str(src), "-frames:v", "1",
                    "-vf", f"scale='min({width},iw)':-2", "-q:v", "3", str(f)], capture_output=True)
    return f.exists()


def _sheets(frames, od, per=20, cols=5, w=320):
    import cv2, numpy as np
    for old in od.glob("sheet*.jpg"):
        old.unlink()
    n = 0
    for k in range(0, len(frames), per):
        imgs = []
        for t, f, *_ in frames[k:k + per]:
            im = cv2.imread(str(f))
            if im is None:
                continue
            im = cv2.resize(im, (w, int(im.shape[0] * w / im.shape[1])))
            hh = im.shape[0]
            cv2.rectangle(im, (0, hh - 22), (92, hh), (0, 0, 0), -1)
            cv2.putText(im, _hms(t), (4, hh - 6), cv2.FONT_HERSHEY_SIMPLEX, 0.5, (255, 255, 255), 1)
            imgs.append(im)
        if not imgs:
            continue
        h = max(i.shape[0] for i in imgs)
        imgs = [cv2.copyMakeBorder(i, 0, h - i.shape[0], 0, 2, cv2.BORDER_CONSTANT) for i in imgs]
        blank = np.zeros_like(imgs[0])
        rows = [np.hstack(imgs[i:i + cols] + [blank] * (cols - len(imgs[i:i + cols])))
                for i in range(0, len(imgs), cols)]
        n += 1
        cv2.imwrite(str(od / f"sheet_{n:02d}.jpg"), np.vstack(rows), [cv2.IMWRITE_JPEG_QUALITY, 80])
    return n


def _dhash(gray, s=16):
    import cv2
    g = cv2.resize(gray, (s + 1, s), interpolation=cv2.INTER_AREA).astype("int16")
    return (g[:, 1:] > g[:, :-1]).flatten()



def _slide_box(live, w, h, min_live=0.02):
    """The 'live' mask (pixels that move most of the time) is the webcam tile. If it sits on the right or at the
    bottom, the slide is everything left of / above it. Returns (x0, y0, x1, y1) in frame pixels or None."""
    import numpy as np
    H, W = live.shape
    if float(live.mean()) < min_live:
        return None
    ys, xs = np.where(live)
    x0, y0 = int(xs.min()), int(ys.min())
    if x0 > W * 0.45:
        return (0, 0, max(200, int(x0 * w / W) - 4), h)
    if y0 > H * 0.45:
        return (0, 0, w, max(120, int(y0 * h / H) - 4))
    return None



def _dedupe(frames, box):
    """Drop frames that look the same in the slide region (same page scrolled a little, presenter moved).
    Keeps the first of each group and deletes the others' files."""
    import cv2
    kept, hashes = [], []
    for fr in frames:
        im = cv2.imread(str(fr[1]), cv2.IMREAD_GRAYSCALE)
        if im is None:
            continue
        if box:
            im = im[box[1]:box[3], box[0]:box[2]]
        hh = _dhash(im)
        if any(int((hh != k).sum()) <= DHASH_MAX_DIST for k in hashes):
            try:
                os.remove(fr[1])
            except OSError:
                pass
            continue
        hashes.append(hh)
        kept.append(fr)
    return kept


_ocr_reader = None



def _get_ocr_reader():
    """media_lab's own EasyOCR reader: GPU (after whisper is done with the file), configurable languages.
    Falls back to CPU if the GPU reader can't be built. Materials/PDF pages keep using reel_lab's CPU reader."""
    global _ocr_reader
    if _ocr_reader is None:
        import logging
        logging.getLogger("easyocr").setLevel(logging.ERROR)
        logging.getLogger("easyocr.easyocr").setLevel(logging.ERROR)
        try:
            import easyocr
            try:
                _ocr_reader = easyocr.Reader(OCR_LANGS, gpu=OCR_GPU)
            except Exception as e:
                print("  note: GPU OCR unavailable, using CPU:", repr(e)[:120], flush=True)
                _ocr_reader = easyocr.Reader(OCR_LANGS, gpu=False)
        except Exception as e:
            print("OCR unavailable:", e, flush=True)
            _ocr_reader = False
    return _ocr_reader



def _ocr(frames, od, box=None):
    """OCR of the slide region of every (deduped) frame. ocr.txt gets one line per change of on-screen text."""
    import cv2
    r = _get_ocr_reader()
    texts = {}
    if not r or not frames:
        return texts
    lines, prev = [], ""
    for t, f, *_ in frames:
        im = cv2.imread(str(f))
        if im is None:
            continue
        if box:
            im = im[box[1]:box[3], box[0]:box[2]]
        try:
            txt = " ".join(x[1] for x in r.readtext(im) if x[2] > 0.4).strip()
        except Exception as e:
            if "out of memory" in str(e).lower():
                import torch
                torch.cuda.empty_cache()
            continue
        texts[t] = txt
        if txt and txt != prev:
            lines.append(f"[{_hms(t)}] {txt}")
            prev = txt
    (od / "ocr.txt").write_text("\n".join(lines))
    return texts


def visuals(src, od, duration, frame_every=None):
    """CPU visual work for one video (runs in parallel with whisper): scan, pick, grab, dedupe, contact sheets.
    OCR is NOT done here any more - process() runs it on the GPU once whisper is finished.
    Returns (frames [(t, file, section_start, reason)], stats, slide_box)."""
    import cv2
    t0 = time.time()
    fd = od / "frames"
    fd.mkdir(parents=True, exist_ok=True)
    for old in fd.glob("*.jpg"):
        old.unlink()
    live = None
    if duration <= SHORT_MAX or frame_every:
        every = frame_every or max(10.0, duration / SHORT_FRAMES)
        times, t = [], every / 2
        while t < duration and len(times) < (SHORT_FRAMES if duration <= SHORT_MAX else LONG_MAX):
            times.append((t, t, "uniform")); t += every
        stats = {"mode": "uniform", "every_s": round(every, 1)}
    else:
        ts, a, mode = _scan(src, duration)
        times, stats = _pick_long(ts, a, duration)
        live = stats.pop("_live", None)
        stats["mode"] = mode
    frames = []
    for i, (t, start, why) in enumerate(times):
        f = fd / f"f{i:03d}_{_hms(t).replace(':', '-')}.jpg"
        if _grab(src, t, f):
            frames.append((t, f, start, why))
    box = None
    if frames and live is not None:
        im = cv2.imread(str(frames[0][1]))
        if im is not None:
            box = _slide_box(live, im.shape[1], im.shape[0])
    n0 = len(frames)
    frames = _dedupe(frames, box)
    stats["frames"] = len(frames)
    stats["dropped_dups"] = n0 - len(frames)
    stats["slide_box"] = list(box) if box else None
    stats["sheets"] = _sheets(frames, od) if frames else 0
    stats["seconds"] = round(time.time() - t0, 1)
    return frames, stats, box


def _outline(od, rel, frames, texts, segs, duration):
    """Timeline for long videos: each visual section with what was said; the on-screen text is printed
    only when it changes (a slide shown for 6 sections is quoted once)."""
    starts = sorted({(fr[2], fr[0]) for fr in frames})
    if not starts:
        starts = [(0.0, None)]
    elif starts[0][0] > 0:
        starts = [(0.0, None)] + starts
    out = [f"# {rel}", f"_{_hms(duration)} · {len(frames)} frames · sections = visual changes · "
           f"'On screen' is quoted only when it changes_", ""]
    prev_screen = None
    for k, (s, tf) in enumerate(starts):
        e = starts[k + 1][0] if k + 1 < len(starts) else duration + 1
        said = " ".join(x["text"] for x in segs if x["start"] is not None and s <= x["start"] < e)
        screen = texts.get(tf, "") if tf is not None else ""
        new_screen = bool(screen) and screen != prev_screen
        if not said and not new_screen:
            continue
        out.append(f"## [{_hms(s)} - {_hms(min(e, duration))}]" + (f" frame {_hms(tf)}" if tf is not None else ""))
        if new_screen:
            out.append(f"**On screen:** {screen}")
        if said:
            out.append(said)
        out.append("")
        if screen:
            prev_screen = screen
    (od / "outline.md").write_text("\n".join(out))


# ---------------- learning materials (CPU only) ----------------
def res_out(folder, src):
    return OUT / Path(folder).name / Path(src).relative_to(folder)


def _vendor():
    if VENDOR.exists() and str(VENDOR) not in sys.path:
        sys.path.append(str(VENDOR))           # LAST: never shadow the env's own PIL / typing_extensions / lxml


def _ocr_image(path):
    r = rl._get_ocr()
    if not r:
        return ""
    try:
        return " ".join(x[1] for x in r.readtext(str(path)) if x[2] > 0.4).strip()
    except Exception:
        return ""


def _text_pdf(src):
    _vendor()
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
    parts, ocrd = [], 0
    with fitz.open(str(src)) as doc:
        for i, page in enumerate(doc, 1):
            t = page.get_text("text").strip()
            if len(t) < 20 and ocrd < PDF_OCR_PAGES:        # scanned page -> OCR on CPU
                img = TMP / f"{os.getpid()}_pdfpage.png"
                page.get_pixmap(dpi=150).save(str(img))
                t = _ocr_image(img)
                img.unlink(missing_ok=True)
                ocrd += 1
            parts.append(f"--- page {i} ---\n{t}")
    return "\n\n".join(parts)


def _text_docx(src):
    _vendor()
    import docx
    d = docx.Document(str(src))
    out = [p.text for p in d.paragraphs]
    for tb in d.tables:
        for row in tb.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(out)


def _text_pptx(src):
    _vendor()
    from pptx import Presentation
    out = []
    for i, slide in enumerate(Presentation(str(src)).slides, 1):
        out.append(f"--- slide {i} ---")
        for sh in slide.shapes:
            if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip():
                out.append(sh.text_frame.text.strip())
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append("[notes] " + slide.notes_slide.notes_text_frame.text.strip())
    return "\n".join(out)


def _text_xlsx(src):
    _vendor()
    import openpyxl
    wb = openpyxl.load_workbook(str(src), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"--- sheet {ws.title} ---")
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                out.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(out)


def _text_html(data):
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(data, "html.parser")
    for t in soup(["script", "style"]):
        t.decompose()
    return re.sub(r"\n{3,}", "\n\n", soup.get_text("\n")).strip()


def _text_epub(src):
    import zipfile
    out = []
    with zipfile.ZipFile(str(src)) as z:
        for n in sorted(z.namelist()):
            if n.lower().endswith((".xhtml", ".html", ".htm")):
                out.append(_text_html(z.read(n)))
    return "\n\n".join(out)


def _unzip(src, dst, info):
    """Course downloads often come zipped: unpack next to the copy and extract text of every document inside."""
    import zipfile
    root = Path(str(dst) + ".unzipped")
    n = chars = 0
    with zipfile.ZipFile(str(src)) as z:
        total = sum(i.file_size for i in z.infolist())
        if total > 4 * RES_MAX:
            info["note"] = f"copied; not unpacked ({total / 1e9:.1f} GB inside)"
            return info
        for i in z.infolist():
            name = i.filename
            if i.is_dir() or name.startswith("__MACOSX") or Path(name).name.startswith("."):
                continue
            target = (root / name).resolve()
            if not str(target).startswith(str(root.resolve())):      # no path traversal
                continue
            target.parent.mkdir(parents=True, exist_ok=True)
            with z.open(i) as fsrc, open(target, "wb") as fdst:
                fdst.write(fsrc.read())
            n += 1
            ext = target.suffix.lower()
            try:
                t = EXTRACT[ext](target) if ext in EXTRACT else (_ocr_image(target) if ext in IMAGES else "")
            except Exception:
                t = ""
            if t.strip():
                Path(str(target) + ".extracted.txt").write_text(t)
                chars += len(t)
    info.update(text_chars=chars, note=f"unpacked {n} files" + (" (media inside zips are not transcribed)" if any(
        Path(x).suffix.lower() in VIDEO | AUDIO for x in zipfile.ZipFile(str(src)).namelist()) else ""))
    return info


EXTRACT = {".pdf": _text_pdf, ".docx": _text_docx, ".pptx": _text_pptx, ".xlsx": _text_xlsx, ".xlsm": _text_xlsx,
           ".html": lambda p: _text_html(Path(p).read_bytes()), ".htm": lambda p: _text_html(Path(p).read_bytes()),
           ".epub": _text_epub}


def _res_done(folder, src):
    dst = res_out(folder, src)
    try:
        a, b = src.stat(), dst.stat()
        return a.st_size == b.st_size and int(a.st_mtime) == int(b.st_mtime)
    except FileNotFoundError:
        return False


def process_resource(folder, src):
    """Copy a learning material 1:1 and write '<name>.extracted.txt' with its text (when possible)."""
    import shutil
    dst = res_out(folder, src)
    dst.parent.mkdir(parents=True, exist_ok=True)
    size = src.stat().st_size
    info = {"rel": str(src.relative_to(folder)), "size": size, "copied": False, "text_chars": 0, "note": ""}
    if size > RES_MAX:
        info["note"] = f"not copied: {size / 1e6:.0f} MB > {RES_MAX / 1e6:.0f} MB limit"
        return info
    shutil.copy2(src, dst)                     # copy2 keeps mtime -> used to skip it next time
    info["copied"] = True
    ext = src.suffix.lower()
    txt = ""
    if ext == ".zip":
        return _unzip(src, dst, info)
    try:
        if ext in EXTRACT:
            txt = EXTRACT[ext](src)
        elif ext in IMAGES:
            txt = _ocr_image(src)
        elif ext in TEXTLIKE:
            info["text_chars"] = size
            return info                        # already text; the copy is enough
        else:
            info["note"] = "copied only (no text extractor for this type)"
    except Exception as e:
        info["note"] = f"copied; text extraction failed: {e!r}"[:200]
    if txt.strip():
        Path(str(dst) + ".extracted.txt").write_text(txt)
        info["text_chars"] = len(txt)
    return info


# ---------------- one file ----------------
def process(folder, src, frame_every=None, ocr=True, pool=None):
    t0 = time.time()
    od = out_dir(folder, src)
    od.mkdir(parents=True, exist_ok=True)
    st = src.stat()
    try:
        info = rl._probe(src)
    except RuntimeError as e:
        if "moov atom not found" in str(e):
            raise Incomplete("file is incomplete (no 'moov' index): still uploading, or the recording was not "
                             "finalised - re-copy / re-export it") from None
        if "End of file" in str(e) or "error reading header" in str(e):
            raise Incomplete("file is incomplete (ends early): still uploading or the copy was interrupted") from None
        raise
    dur = info["duration"]
    is_video = src.suffix.lower() in VIDEO and bool(info.get("width"))
    meta = {"path": str(src), "rel": str(src.relative_to(folder)), "size": st.st_size, "mtime": int(st.st_mtime),
            "audio_seconds": None,
            "kind": "video" if is_video else "audio", "duration": round(dur, 1), "duration_hms": _hms(dur),
            "width": info.get("width"), "height": info.get("height"), "ok": False}
    x = None
    if info["has_audio"]:
        x = _load_audio(src)
        meta["audio_seconds"] = round(_check_complete(x, dur), 1)     # reject partial files before any work
    vis = pool.submit(visuals, src, od, dur, frame_every) if (is_video and pool) else None
    segs, text, asr_s = [], "", 0.0
    if x is not None:
        ta = time.time()
        segs = _transcribe(x)
        asr_s = time.time() - ta
        text = _paragraphs(segs)
        del x
    (od / "segments.json").write_text(json.dumps(segs, ensure_ascii=False, indent=0))
    (od / "transcript.txt").write_text(text)
    (od / "transcript_ts.txt").write_text("\n".join(f"[{_hms(s['start'])}] {s['text']}" for s in segs))
    words = len(text.split())
    meta.update(words=words, segments=len(segs), wpm=round(words / (dur / 60), 1) if dur else None,
                asr_seconds=round(asr_s, 1), asr_speed=round(dur / asr_s, 1) if asr_s else None)
    if is_video:
        frames, vstats, box = vis.result() if vis else visuals(src, od, dur, frame_every)
        to = time.time()
        try:
            import torch
            torch.cuda.empty_cache()  # whisper's activations are gone; OCR gets the head-room on the GPU
        except Exception:
            pass
        texts = _ocr(frames, od, box) if ocr else {}
        vstats["ocr_lines"] = sum(1 for v in texts.values() if v)
        vstats["ocr_seconds"] = round(time.time() - to, 1)
        meta["visual"] = vstats
        meta["frames"] = vstats.get("frames")
        if dur > SHORT_MAX and frames:
            _outline(od, meta["rel"], frames, texts, segs, dur)
    meta.update(ok=True, processed_at=time.strftime("%Y-%m-%d %H:%M:%S"), seconds=round(time.time() - t0, 1))
    (od / "meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=1))
    (od / "error.txt").unlink(missing_ok=True)
    with open(INDEX, "a") as f:
        f.write(json.dumps({k: meta.get(k) for k in ("path", "duration", "words", "frames", "processed_at")},
                           ensure_ascii=False) + "\n")
    return meta


GENERATED = ("_all_transcripts.md", "_all_materials.md", "_manifest.md")


def _combine(folder):
    """Build the folder summaries from the OUTPUT folder only (works after the source folder is deleted)."""
    base = OUT / Path(folder).name
    if not base.exists():
        return
    media_dirs = {m.parent for m in base.rglob("meta.json")} | {e.parent for e in base.rglob("error.txt")}
    trans, mats = [], []
    man = ["# Manifest: " + base.name, "", f"_generated {time.strftime('%Y-%m-%d %H:%M')}_", "",
           "| File | Kind | Length / size | Text | Status |", "|---|---|---|---|---|"]
    for d in sorted(media_dirs):
        rel = d.relative_to(base)
        m = d / "meta.json"
        if m.exists():
            x = json.loads(m.read_text())
            src_t = d / "outline.md" if (d / "outline.md").exists() else d / "transcript.txt"
            t = src_t.read_text() if src_t.exists() else ""
            trans.append(f"# {x['rel']}\n_{x['kind']} · {x['duration_hms']} · {x.get('words', 0)} words · "
                         f"{x.get('wpm')} wpm_\n\n{t}\n")
            extra = " + outline" if (d / "outline.md").exists() else ""
            man.append(f"| {rel} | {x['kind']} | {x['duration_hms']} | {x.get('words', 0)} words{extra} | ok |")
        else:
            why = (d / "error.txt").read_text()[:120].replace("|", "/").replace("\n", " ")
            man.append(f"| {rel} | media | - | - | FAILED: {why} |")
    for f in sorted(base.rglob("*")):
        if f.is_dir() or f.name in GENERATED or f.name.endswith(".extracted.txt"):
            continue
        if any(p in media_dirs for p in f.parents):
            continue
        rel = f.relative_to(base)
        side = Path(str(f) + ".extracted.txt")
        text = ""
        if side.exists():
            text = side.read_text(errors="ignore")
        elif f.suffix.lower() in TEXTLIKE and f.stat().st_size < 5_000_000:
            text = f.read_text(errors="ignore")
        man.append(f"| {rel} | material ({f.suffix.lower() or 'no ext'}) | {f.stat().st_size / 1e6:.1f} MB | "
                   f"{len(text):,} chars | {'text' if text.strip() else 'copied'} |")
        if text.strip():
            mats.append(f"# {rel}\n\n{text.strip()}\n")
    (base / "_all_transcripts.md").write_text("\n\n".join(trans))
    (base / "_all_materials.md").write_text("\n\n".join(mats))
    (base / "_manifest.md").write_text("\n".join(man) + "\n")
    try:
        z = _zip(base)
        print(time.strftime("%H:%M:%S"), "zip", z.name, f"({z.stat().st_size / 1e6:.0f} MB)", flush=True)
    except Exception as e:
        print("zip failed:", repr(e)[:200], flush=True)


def _zip(base):
    """out/<folder>.zip = the whole output folder, rebuilt after every run (old zip deleted first).
    Written as .zip.part and renamed at the end, so a half-written zip is never mistaken for a finished one."""
    import zipfile
    z = base.parent / (base.name + ".zip")
    part = base.parent / (base.name + ".zip.part")
    for old in (z, part):
        try:
            old.unlink()
        except FileNotFoundError:
            pass
    with zipfile.ZipFile(str(part), "w", zipfile.ZIP_DEFLATED, compresslevel=6) as zf:
        for f in sorted(base.rglob("*")):
            if f.is_file():
                zf.write(str(f), str(f.relative_to(base.parent)))
    os.replace(str(part), str(z))
    return z


# ---------------- worker ----------------
def _acquire_lock():
    while True:
        if rl.LOCK.exists():
            try:
                pid = int(rl.LOCK.read_text().strip() or 0)
            except ValueError:
                pid = 0
            if pid and pid != os.getpid() and _alive(pid):
                _state(status=f"waiting for GPU lock (pid {pid})")
                time.sleep(15)
                continue
        rl.LOCK.write_text(str(os.getpid()))
        return


def _cleanup():
    try:
        rl.release()
    except Exception:
        pass
    for f in glob.glob(str(TMP / "*")):
        try:
            os.remove(f)
        except OSError:
            pass
    try:
        if rl.LOCK.exists() and rl.LOCK.read_text().strip() == str(os.getpid()):
            rl.LOCK.unlink(missing_ok=True)
    except Exception:
        pass


def _release_gpu():
    try:
        rl.release()
    except Exception:
        pass
    try:
        if rl.LOCK.exists() and rl.LOCK.read_text().strip() == str(os.getpid()):
            rl.LOCK.unlink(missing_ok=True)
    except Exception:
        pass


def _job(job_file):
    from concurrent.futures import ThreadPoolExecutor
    job = json.loads(Path(job_file).read_text())

    def _on_term(*_):
        raise SystemExit("stopped")
    signal.signal(signal.SIGTERM, _on_term)
    try:
        os.nice(10)                                   # production comes first
    except OSError:
        pass
    try:
        import torch
        torch.set_num_threads(CPU_THREADS)
    except Exception:
        pass
    folder = Path(job["folder"])
    force = bool(job.get("force"))
    status = "finished"
    pool = ThreadPoolExecutor(max_workers=1)
    new, prev, failures = set(), set(), {}          # media: done this run / done before / failed {rel: {...}}
    mat_new, mat_prev, mat_fail = set(), set(), {}
    durs, waiting = {}, []
    t_start, media_done, wait_since = time.time(), 0.0, None
    have_lock = False
    print(f"=== RUN {time.strftime('%Y-%m-%d %H:%M:%S')} | {folder} ===", flush=True)

    def dur_of(src):
        sig = (str(src), src.stat().st_size, int(src.stat().st_mtime))
        if sig not in durs:
            durs[sig] = _duration(src)
        return durs[sig]

    def publish(**kw):
        try:
            _, media, res = _walk(folder)
        except FileNotFoundError:
            media, res = [], []
        todo_s = 0.0
        for m in media:
            r = str(m.relative_to(folder))
            if r in new or r in prev or r in failures or _uploading(m):
                continue
            try:
                todo_s += dur_of(m)
            except FileNotFoundError:
                pass
        el = time.time() - t_start
        speed = media_done / el if el > 60 and media_done else None
        _state(pid=os.getpid(), folder=str(folder), status=kw.pop("status", "running"),
               media_total=len(media), media_new=len(new), media_prev=len(prev), media_failed=len(failures),
               materials_total=len(res), materials_new=len(mat_new), materials_prev=len(mat_prev),
               materials_failed=len(mat_fail), waiting=waiting,
               failures={k: v["why"] for k, v in failures.items()} | {k: v for k, v in mat_fail.items()},
               hours_done=round(media_done / 3600, 2), hours_left=round(todo_s / 3600, 2),
               speed_x=round(speed, 1) if speed else None,
               eta=f"{todo_s / speed / 3600:.1f} h" if speed else ("0 h" if not todo_s else "calculating"),
               **kw)

    try:
        _state(pid=os.getpid(), folder=str(folder), status="starting", started=time.strftime("%Y-%m-%d %H:%M:%S"),
               finished=None, current=None, run_marker=f"=== RUN")
        while True:
            _, media, res = _walk(folder)
            waiting = []
            # 1) learning materials: CPU only, quick
            for src in res:
                rel = str(src.relative_to(folder))
                if rel in mat_new or rel in mat_prev:
                    continue
                if _uploading(src):
                    waiting.append(rel); continue
                if not force and _res_done(folder, src):
                    mat_prev.add(rel); continue
                publish(current=f"material: {rel}", current_since=time.strftime("%H:%M:%S"))
                try:
                    info = process_resource(folder, src)
                    mat_new.add(rel); mat_fail.pop(rel, None)
                    print(time.strftime("%H:%M:%S"), "material", rel, f"({info['text_chars']:,} chars of text)"
                          + (f" - {info['note']}" if info["note"] else ""), flush=True)
                except Exception as e:
                    mat_fail[rel] = f"material: {e!r}"[:200]
                    print(time.strftime("%H:%M:%S"), "material FAILED", rel, repr(e)[:200], flush=True)
            # 2) media: GPU
            for src in media:
                rel = str(src.relative_to(folder))
                if rel in new or rel in prev:
                    continue
                if not force and _is_done(folder, src):
                    prev.add(rel); continue
                if _uploading(src):
                    waiting.append(rel); continue
                sig = (src.stat().st_size, int(src.stat().st_mtime))
                if rel in failures and failures[rel]["sig"] == sig:
                    continue                                  # unchanged since it failed: don't retry
                if not have_lock:
                    publish(status="waiting for GPU lock")
                    _acquire_lock(); have_lock = True
                d = dur_of(src)
                publish(current=f"{rel} ({_hms(d)})", current_since=time.strftime("%H:%M:%S"))
                print(time.strftime("%H:%M:%S"), "media", rel, _hms(d), flush=True)
                try:
                    m = process(folder, src, job.get("frame_every"), job.get("ocr", True), pool)
                    new.add(rel); failures.pop(rel, None); media_done += d
                    v = m.get("visual") or {}
                    print(f"   ok {m['words']}w, asr {m.get('asr_speed')}x, frames {v.get('frames', '-')} "
                          f"({v.get('mode', 'audio')}), {m['seconds']}s", flush=True)
                except SystemExit:
                    raise
                except Exception as e:
                    why = str(e) if isinstance(e, Incomplete) else repr(e)
                    failures[rel] = {"why": why[:240], "sig": sig}
                    print("   FAILED", why[:300], flush=True)
                    od = out_dir(folder, src)
                    od.mkdir(parents=True, exist_ok=True)
                    (od / "error.txt").write_text(why)
                try:
                    import torch
                    torch.cuda.empty_cache()
                except Exception:
                    pass
            _combine(folder)
            if not waiting or not job.get("wait", True):
                break
            # uploads still in progress: free the GPU and rescan every minute
            if have_lock:
                _release_gpu(); have_lock = False
            wait_since = wait_since or time.time()
            if time.time() - wait_since > WAIT_MAX_S:
                print("gave up waiting for uploads:", ", ".join(waiting), flush=True)
                break
            publish(status=f"waiting for {len(waiting)} file(s) still uploading", current=None)
            time.sleep(60)
            if not any(_uploading(folder / w) for w in waiting if (folder / w).exists()):
                wait_since = None                              # uploads finished; next pass processes them
    except SystemExit:
        status = "stopped"
    except Exception as e:
        status = f"crashed: {e!r}"[:300]
        print(status, flush=True)
    finally:
        pool.shutdown(wait=False, cancel_futures=True)
        _cleanup()
        try:
            publish(status=status, current=None, finished=time.strftime("%Y-%m-%d %H:%M:%S"))
        except Exception:
            _state(status=status, current=None, finished=time.strftime("%Y-%m-%d %H:%M:%S"))
        try:
            os.remove(job_file)
        except OSError:
            pass
        print(f"DONE {status}: media {len(new)} new / {len(prev)} already done / {len(failures)} failed; "
              f"materials {len(mat_new)} new / {len(mat_prev)} already copied / {len(mat_fail)} failed", flush=True)


# ---------------- notebook API ----------------
def run_folder(folder, frame_every=None, ocr=True, force=False, wait=True):
    """Start processing FOLDER in the background. Returns immediately; check status().
    wait=True: files still uploading are picked up when their upload finishes (the worker keeps rescanning).
    frame_every: force uniform frames every N seconds (default: smart). force: redo everything."""
    folder, media, res = _walk(folder)
    st = _read_state()
    if st.get("pid") and _alive(st["pid"]) and _is_our_worker(st["pid"]):
        raise RuntimeError(f"media_lab is already running (pid {st['pid']}) on {st.get('folder')} - wait or stop()")
    job = JOBS / f"{int(time.time())}.json"
    job.write_text(json.dumps({"folder": str(folder), "frame_every": frame_every, "ocr": ocr, "force": force,
                               "wait": wait}))
    env = dict(os.environ, CUDA_DEVICE_ORDER="PCI_BUS_ID", CUDA_VISIBLE_DEVICES=rl._gpu_uuid("4090"), REEL_LAB_GPU="0",
               PYTHONWARNINGS="ignore", TRANSFORMERS_VERBOSITY="error", OMP_NUM_THREADS=str(CPU_THREADS),
               MKL_NUM_THREADS=str(CPU_THREADS))
    log = LOGS / f"{folder.name}.log"
    with open(log, "a") as lf:
        proc = subprocess.Popen([sys.executable, str(HERE / "media_lab.py"), "--job", str(job)], env=env, cwd=str(HERE),
                                stdout=lf, stderr=subprocess.STDOUT, start_new_session=True)
    STATE.write_text(json.dumps({"pid": proc.pid, "folder": str(folder), "status": "starting",
                                 "media_total": len(media), "materials_total": len(res), "log": str(log),
                                 "started": time.strftime("%Y-%m-%d %H:%M:%S"),
                                 "updated": time.strftime("%Y-%m-%d %H:%M:%S")}, indent=1))
    return (f"started pid {proc.pid}: {len(media)} media + {len(res)} learning materials in {folder}\n"
            f"check progress with ml.status()")


def _read_state():
    try:
        return json.loads(STATE.read_text())
    except Exception:
        return {}


def _ago(ts):
    try:
        t = time.mktime(time.strptime(ts, "%Y-%m-%d %H:%M:%S"))
        return _hms(time.time() - t)
    except Exception:
        return "?"


def status(tail=15):
    """Full dashboard of the current / last run."""
    st = _read_state()
    if not st:
        print("idle (never run)"); return
    alive = bool(st.get("pid") and _alive(st["pid"]) and _is_our_worker(st["pid"]))
    s = st.get("status", "?")
    head = ("RUNNING" if s == "running" else s.upper()) if alive else \
        ("FINISHED" if s == "finished" else f"NOT RUNNING (last status: {s})")
    L = [f"{'●' if alive else '○'} {head}",
         f"folder   : {st.get('folder')}",
         f"started  : {st.get('started')}   (running for {_ago(st['started']) if alive and st.get('started') else '-'})"
         + (f"   finished: {st['finished']}" if st.get("finished") else "")]
    L.append(f"media    : {st.get('media_new', 0)} done this run · {st.get('media_prev', 0)} already done earlier · "
             f"{st.get('media_failed', 0)} failed · {sum(1 for w in st.get('waiting', []) if Path(w).suffix.lower() in VIDEO | AUDIO)} "
             f"uploading · of {st.get('media_total', '?')}")
    L.append(f"materials: {st.get('materials_new', 0)} copied this run · {st.get('materials_prev', 0)} already copied · "
             f"{st.get('materials_failed', 0)} failed · of {st.get('materials_total', '?')}")
    L.append(f"hours    : {st.get('hours_done', 0)} h transcribed this run · {st.get('hours_left', '?')} h left · "
             f"speed {st.get('speed_x') or '-'}x · ETA {st.get('eta', '-')}")
    if alive and st.get("current"):
        L.append(f"now      : {st['current']}   (since {st.get('current_since', '?')})")
    fails = st.get("failures") or {}
    if fails:
        L.append(f"\nFAILED ({len(fails)}) - re-copy the file: a changed file is retried automatically (or on the next run_folder):")
        L += [f"  ✗ {k}\n      {v}" for k, v in fails.items()]
    wait = st.get("waiting") or []
    if wait:
        L.append(f"\nWAITING FOR UPLOAD ({len(wait)}) - picked up automatically when the upload finishes:")
        L += [f"  … {w}" for w in wait]
    if st.get("log") and Path(st["log"]).exists():
        lines = Path(st["log"]).read_text(errors="ignore").splitlines()
        k = max((i for i, l in enumerate(lines) if l.startswith("=== RUN")), default=0)
        run = [l for l in lines[k:] if l.strip()]
        L.append(f"\n--- log of this run (last {min(tail, len(run))} of {len(run)} lines; ml.log() for all) ---")
        L += run[-tail:]
    print("\n".join(L))


def log(lines=200):
    """Full log of the current / last run."""
    st = _read_state()
    if st.get("log") and Path(st["log"]).exists():
        all_ = Path(st["log"]).read_text(errors="ignore").splitlines()
        k = max((i for i, l in enumerate(all_) if l.startswith("=== RUN")), default=0)
        print("\n".join(all_[k:][-lines:]))


def stop():
    """Stop OUR worker only (SIGTERM -> it frees GPU + tmp itself)."""
    st = _read_state()
    pid = st.get("pid")
    if pid and _alive(pid) and _is_our_worker(pid):
        os.kill(int(pid), signal.SIGTERM)
        return f"sent stop to pid {pid}"
    return "nothing running"


def overview(folder):
    import pandas as pd
    base = OUT / Path(folder).expanduser().resolve().name
    rows = []
    for m in sorted(base.rglob("meta.json")):
        d = json.loads(m.read_text())
        v = d.pop("visual", None) or {}
        d["visual_mode"] = v.get("mode")
        rows.append(d)
    mats = [f for f in base.rglob("*") if f.is_file() and f.name not in GENERATED
            and not f.name.endswith(".extracted.txt")
            and not any((p / "meta.json").exists() or (p / "error.txt").exists() for p in f.parents)]
    texted = sum(1 for f in mats if Path(str(f) + ".extracted.txt").exists() or f.suffix.lower() in TEXTLIKE)
    print(f"{len(rows)} media files · {len(mats)} learning materials ({texted} with text) · "
          f"see {base}/_manifest.md, _all_transcripts.md, _all_materials.md")
    if not rows:
        return pd.DataFrame()
    df = pd.DataFrame(rows)
    cols = [c for c in ("rel", "kind", "duration_hms", "words", "wpm", "frames", "visual_mode", "asr_speed", "seconds")
            if c in df]
    return df[cols]


def read(folder, rel, what="transcript.txt"):
    """Read one output, e.g. read(FOLDER, 'week1/call.mp4') or read(FOLDER, 'week1/call.mp4', 'outline.md')."""
    f = Path(folder).expanduser().resolve()
    return (out_dir(f, f / rel) / what).read_text()


if __name__ == "__main__" and sys.argv[1:2] == ["--job"]:
    _job(sys.argv[2])
